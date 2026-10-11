"""Classeur « À valider avec Nathalie » : un onglet par sujet, avec une
colonne vide à remplir. Aucun accès réseau.

Usage : python3 -I outils/classeur_nathalie.py inventaire arret-1/a-valider-avec-nathalie.xlsx
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

A_REMPLIR = PatternFill("solid", fgColor="FFF4CC")
ENTETE = Font(bold=True)


def lire(path):
    return list(csv.DictReader(open(path, encoding="utf-8")))


def onglet(wb, titre, explication, colonnes, lignes, a_remplir, largeurs):
    ws = wb.create_sheet(titre)
    ws["A1"] = explication
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(colonnes) + len(a_remplir))
    ws.row_dimensions[1].height = 60
    entetes = [c[0] for c in colonnes] + a_remplir
    for j, e in enumerate(entetes, 1):
        c = ws.cell(row=2, column=j, value=e)
        c.font = ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if e in a_remplir:
            c.fill = A_REMPLIR
    for i, l in enumerate(lignes, 3):
        for j, (_, cle) in enumerate(colonnes, 1):
            v = l.get(cle, "") if isinstance(cle, str) else cle(l)
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for k in range(len(a_remplir)):
            ws.cell(row=i, column=len(colonnes) + 1 + k).fill = A_REMPLIR
    for j, w in enumerate(largeurs, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A3"


def main(inv, dest):
    inv = Path(inv)
    wb = Workbook()
    lisez = wb.active
    lisez.title = "Lisez-moi"
    texte = [
        "À valider avec Nathalie — classeur préparé le 11 octobre 2026",
        "",
        "Chaque onglet contient ce que j'ai lu sur le site actuel, et une ou plusieurs colonnes jaunes à remplir.",
        "Rien n'a été deviné d'après les photos : tout vient du texte des fiches et des pages.",
        "Une case qu'on ne sait pas remplir : écrire « je ne sais pas ». Une proposition qui convient : écrire « ok ».",
        "",
        "Onglets :",
        "1. Titres — titre actuel, titre proposé d'après la description, ce qui reste à vérifier.",
        "2. Designers — noms cités sur les fiches, et le niveau d'attribution (de, attribué à, dans le goût de).",
        "3. Douane et envoi — type douanier et taille d'envoi proposés pour chaque pièce, à valider.",
        "4. Projets — les pages actuelles que je propose de réunir en une seule.",
        "5. Presse — les dates lues qui semblent incohérentes.",
        "6. Photos — les pièces à rephotographier, les plus chères d'abord.",
        "",
        "Aucun prix en dollars ne paraîtra sans son accord sur le prix export, le type douanier et la taille d'envoi.",
    ]
    for i, t in enumerate(texte, 1):
        lisez.cell(row=i, column=1, value=t)
    lisez["A1"].font = Font(bold=True, size=14)
    lisez.column_dimensions["A"].width = 120

    pieces = {p["id_boutique"]: p for p in lire(inv / "pieces.csv")}
    onglet(wb, "1. Titres",
           "Titres trop vagues ou en double. Le titre proposé reprend la description de la fiche. Colonne jaune : le titre à garder "
           "(« ok » si la proposition convient), et les informations manquantes si elle les connaît (époque, designer ou fabricant, matières).",
           [("Réf.", "id_boutique"), ("Titre actuel", "titre_actuel"), ("Titre proposé", "titre_propose"),
            ("Ce que dit la fiche", "ce_que_dit_la_description"), ("À vérifier", "a_verifier")],
           lire(inv / "titres_a_valider.csv"), ["Titre validé", "Époque, designer, matières (si connus)"],
           [8, 28, 40, 60, 28, 40, 40])

    onglet(wb, "2. Designers",
           "Noms cités sur les fiches actuelles (designers, ateliers, éditeurs, artistes, maisons de tissus). Le site ne citera que ceux-là. "
           "Colonne jaune : le niveau exact pour chaque pièce (de / attribué à / dans le goût de), l'orthographe à retenir, et si le nom doit "
           "apparaître dans les filtres. Question à part : pour quoi la galerie est-elle connue ?",
           [("Nom tel qu'écrit", "nom_tel_qu_ecrit"), ("Nb de pièces", "nb_pieces"), ("Pièces", "pieces"),
            ("Niveau lu dans le texte", "niveau_lu_dans_le_texte"), ("Remarque", "remarque")],
           lire(inv / "designers_nommes.csv"), ["Orthographe et niveau validés", "Dans les filtres ? (oui/non)"],
           [24, 8, 60, 26, 30, 36, 18])

    douane = [p for p in pieces.values()]
    onglet(wb, "3. Douane et envoi",
           "Proposé d'après le titre, la description et les dimensions, jamais d'après les photos. Vide = le texte ne suffit pas. "
           "Taux vérifiés le 10/10/2026 pour des pièces fabriquées dans l'UE. Sans validation, la pièce n'aura pas de prix en dollars.",
           [("Réf.", "id_boutique"), ("Titre actuel", "titre_actuel"), ("Famille", "famille"),
            ("Prix affiché TTC (€)", "prix_affiche_ttc"), ("Plus grande dimension (cm)", "plus_grande_dimension_cm"),
            ("Origine lue", "origine_lue"), ("Essences lues", "essences_lues"),
            ("Type douanier proposé", "type_douanier_propose"), ("Remarque", "note_type"),
            ("Taille d'envoi proposée", "taille_envoi_proposee"), ("Remarque taille", "note_taille")],
           douane, ["Type validé", "Taille validée", "Fabriquée dans l'UE ? (oui/non/je ne sais pas)", "Essences de bois",
                    "Prix export hors TVA (€)"],
           [8, 28, 14, 10, 10, 16, 16, 26, 30, 12, 18, 20, 12, 18, 18, 14])

    projets = [p for p in lire(inv / "projets.csv") if p.get("regroupement_propose") and p["rubrique_actuelle"] != "Graphisme"]
    onglet(wb, "4. Projets",
           "« Un projet, une page » : les pages actuelles que je propose de réunir. Colonne jaune : ok / non, et l'année du projet si elle est "
           "connue avec certitude (sinon, rien ne s'affiche).",
           [("Nouvelle page proposée", "regroupement_propose"), ("Page actuelle", "titre_actuel"), ("Adresse actuelle", "ancienne_adresse"),
            ("Lieu écrit", "lieu_ecrit_sur_la_page"), ("Photos", "nb_photos"), ("Crédit écrit", "credits_ecrits_sur_la_page")],
           projets, ["Réunir ? (ok / non)", "Année (si sûre)"], [26, 40, 50, 30, 8, 30, 16, 14])

    presse = [p for p in lire(inv / "presse.csv") if p.get("controle_date")]
    onglet(wb, "5. Presse",
           "Dates lues dans le nom du fichier mais qui ne collent pas avec la date de mise en ligne. Colonne jaune : la date de parution exacte.",
           [("Média (titre de la vignette)", "media_titre_wp"), ("Fichier", lambda l: l["cible"].rsplit("/", 1)[-1]),
            ("Date lue", "date_lisible"), ("Mis en ligne le", "mis_en_ligne_wp"), ("Pourquoi vérifier", "controle_date")],
           presse, ["Date de parution exacte"], [26, 50, 10, 14, 36, 22])

    photos = lire(inv / "a_rephotographier.csv")
    onglet(wb, "6. Photos",
           "Cible : au moins 5 photos de 2 000 px sur le grand côté. Classé du plus cher au moins cher. Colonne jaune : existe-t-il des "
           "photos d'origine plus grandes ? (oui : où ? / non : à refaire)",
           [("Rang", "rang"), ("Réf.", "id_boutique"), ("Titre actuel", "titre_actuel"), ("Prix TTC (€)", "prix_affiche_ttc"),
            ("Nb de photos", "nb_photos"), ("Plus grand côté (px)", "plus_grand_cote_px_max"), ("Ce qui manque", "manque")],
           photos, ["Originaux plus grands ? Où ?"], [6, 8, 30, 10, 8, 10, 44, 30])

    Path(dest).parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    print("écrit :", dest)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
